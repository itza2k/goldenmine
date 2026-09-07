from __future__ import annotations

from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from goldmine.db.connection import Database, rows_to_dicts
from goldmine.exceptions import ValidationError
from goldmine.paths import exports_dir
from goldmine.security.permissions import CurrentUser, require_owner
from goldmine.services.audit import AuditService
from goldmine.util import money, now


class ReportService:
    def __init__(self, db: Database, audit: AuditService, settings) -> None:
        self.db = db
        self.audit = audit
        self.settings = settings

    def dashboard(self) -> dict:
        today = now().date().isoformat()
        due_days = self.settings.int_value("due_soon_days", 7)
        until = (now().date() + timedelta(days=due_days)).isoformat()
        active = self.db.fetchone("SELECT COUNT(*) AS c, COALESCE(SUM(loan_amount),0) AS s FROM loans WHERE status='open'")
        issued = self.db.fetchone("SELECT COUNT(*) AS c FROM loans WHERE date(created_at) = ?", (today,))
        closed = self.db.fetchone("SELECT COUNT(*) AS c FROM loans WHERE status='closed' AND closing_date = ?", (today,))
        collections = self.db.fetchone(
            "SELECT COALESCE(SUM(total_received),0) AS s FROM loans WHERE status='closed' AND closing_date = ?",
            (today,),
        )
        due = self.db.fetchall(
            """
            SELECT l.loan_number, l.due_date, l.loan_amount, c.name AS customer_name
            FROM loans l JOIN customers c ON c.id = l.customer_id
            WHERE l.status='open' AND l.due_date IS NOT NULL AND l.due_date >= ? AND l.due_date <= ?
            ORDER BY l.due_date LIMIT 10
            """,
            (today, until),
        )
        recent = self.db.fetchall("SELECT id, name, phone, created_at FROM customers ORDER BY id DESC LIMIT 8")
        return {
            "active_loans": active["c"],
            "outstanding": active["s"],
            "issued_today": issued["c"],
            "closed_today": closed["c"],
            "collections_today": collections["s"],
            "due_soon": rows_to_dicts(due),
            "recent_customers": rows_to_dicts(recent),
            "due_soon_days": due_days,
        }

    def _period(self, kind: str, date_from: str | None, date_to: str | None) -> tuple[str, str, str]:
        today = now().date()
        if kind == "daily":
            d = date_from or today.isoformat()
            return d, d, f"Daily report — {d}"
        if kind == "weekly":
            start = today - timedelta(days=today.weekday())
            end = start + timedelta(days=6)
            a = date_from or start.isoformat()
            b = date_to or end.isoformat()
            return a, b, f"Weekly report — {a} to {b}"
        if kind == "monthly":
            start = today.replace(day=1)
            a = date_from or start.isoformat()
            b = date_to or today.isoformat()
            return a, b, f"Monthly report — {a} to {b}"
        if kind == "yearly":
            start = today.replace(month=1, day=1)
            a = date_from or start.isoformat()
            b = date_to or today.isoformat()
            return a, b, f"Yearly report — {a} to {b}"
        if kind == "open":
            return "0000-01-01", "9999-12-31", "Open loans"
        if kind == "closed":
            return "0000-01-01", "9999-12-31", "Closed loans"
        if kind == "interest":
            a = date_from or (today.replace(day=1)).isoformat()
            b = date_to or today.isoformat()
            return a, b, f"Interest earned — {a} to {b}"
        if kind == "activity":
            a = date_from or today.isoformat()
            b = date_to or today.isoformat()
            return a, b, f"Employee activity — {a} to {b}"
        raise ValidationError("Unknown report type.")

    def generate(self, user: CurrentUser, kind: str, date_from: str | None = None, date_to: str | None = None) -> dict:
        require_owner(user)
        start, end, title = self._period(kind, date_from, date_to)
        symbol = self.settings.get("currency_symbol", "₹")
        if kind == "activity":
            rows = rows_to_dicts(
                self.db.fetchall(
                    """
                    SELECT created_at, username, action, entity_type, entity_id, reason
                    FROM audit_log
                    WHERE date(created_at) >= ? AND date(created_at) <= ?
                      AND username IN (SELECT username FROM users WHERE role = 'employee')
                    ORDER BY id DESC
                    """,
                    (start, end),
                )
            )
            columns = ["created_at", "username", "action", "entity_type", "entity_id", "reason"]
            headers = ["When", "User", "Action", "Type", "ID", "Reason"]
            summary = {"events": len(rows)}
        elif kind == "open":
            rows = rows_to_dicts(
                self.db.fetchall(
                    """
                    SELECT l.loan_number, c.name, c.phone, l.gold_description, l.gold_weight,
                           l.loan_amount, l.interest_rate, l.start_date, l.due_date
                    FROM loans l JOIN customers c ON c.id = l.customer_id
                    WHERE l.status = 'open'
                    ORDER BY l.start_date
                    """
                )
            )
            columns = [
                "loan_number", "name", "phone", "gold_description", "gold_weight",
                "loan_amount", "interest_rate", "start_date", "due_date",
            ]
            headers = ["Loan #", "Customer", "Phone", "Gold", "Weight", "Amount", "Rate %", "Start", "Due"]
            total = sum(r["loan_amount"] or 0 for r in rows)
            summary = {"count": len(rows), "outstanding": total}
        elif kind == "closed":
            rows = rows_to_dicts(
                self.db.fetchall(
                    """
                    SELECT l.loan_number, c.name, l.loan_amount, l.interest_collected, l.total_received,
                           l.start_date, l.closing_date
                    FROM loans l JOIN customers c ON c.id = l.customer_id
                    WHERE l.status = 'closed'
                    ORDER BY l.closing_date DESC
                    """
                )
            )
            columns = [
                "loan_number", "name", "loan_amount", "interest_collected", "total_received",
                "start_date", "closing_date",
            ]
            headers = ["Loan #", "Customer", "Principal", "Interest", "Received", "Start", "Closed"]
            summary = {
                "count": len(rows),
                "interest": sum(r["interest_collected"] or 0 for r in rows),
                "received": sum(r["total_received"] or 0 for r in rows),
            }
        elif kind == "interest":
            rows = rows_to_dicts(
                self.db.fetchall(
                    """
                    SELECT l.loan_number, c.name, l.closing_date, l.interest_collected, l.total_received
                    FROM loans l JOIN customers c ON c.id = l.customer_id
                    WHERE l.status = 'closed' AND l.closing_date >= ? AND l.closing_date <= ?
                    ORDER BY l.closing_date
                    """,
                    (start, end),
                )
            )
            columns = ["loan_number", "name", "closing_date", "interest_collected", "total_received"]
            headers = ["Loan #", "Customer", "Closed", "Interest", "Received"]
            summary = {
                "count": len(rows),
                "interest": sum(r["interest_collected"] or 0 for r in rows),
            }
        else:
            rows = rows_to_dicts(
                self.db.fetchall(
                    """
                    SELECT l.loan_number, c.name, c.phone, l.gold_description, l.loan_amount,
                           l.interest_rate, l.start_date, l.status, l.total_received
                    FROM loans l JOIN customers c ON c.id = l.customer_id
                    WHERE date(l.created_at) >= ? AND date(l.created_at) <= ?
                    ORDER BY l.id
                    """,
                    (start, end),
                )
            )
            columns = [
                "loan_number", "name", "phone", "gold_description", "loan_amount",
                "interest_rate", "start_date", "status", "total_received",
            ]
            headers = ["Loan #", "Customer", "Phone", "Gold", "Amount", "Rate %", "Start", "Status", "Received"]
            summary = {
                "count": len(rows),
                "issued": sum(r["loan_amount"] or 0 for r in rows),
            }
        self.audit.record(user, "reports.view", entity_type="report", new={"kind": kind, "title": title})
        return {
            "kind": kind,
            "title": title,
            "start": start,
            "end": end,
            "rows": rows,
            "columns": columns,
            "headers": headers,
            "summary": summary,
            "symbol": symbol,
        }

    def export_excel(self, user: CurrentUser, report: dict) -> Path:
        require_owner(user)
        path = exports_dir() / f"{_slug(report['title'])}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        gold = PatternFill("solid", fgColor="C9A227")
        header_font = Font(bold=True, color="1B2A4A")
        ws.append([report["title"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Generated {now().strftime('%Y-%m-%d %H:%M')}"])
        ws.append([])
        ws.append(report["headers"])
        for col, _ in enumerate(report["headers"], 1):
            cell = ws.cell(4, col)
            cell.fill = gold
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for row in report["rows"]:
            ws.append([row.get(c) for c in report["columns"]])
        ws.append([])
        for k, v in report["summary"].items():
            ws.append([k, v])
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(width, 40)
        wb.save(path)
        self.audit.record(user, "reports.export", entity_type="report", new={"format": "excel", "path": path.name})
        return path

    def export_pdf(self, user: CurrentUser, report: dict) -> Path:
        require_owner(user)
        path = exports_dir() / f"{_slug(report['title'])}.pdf"
        doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=40, bottomMargin=36)
        styles = getSampleStyleSheet()
        story = [
            Paragraph(self.settings.get("shop_name", "Goldmine"), styles["Title"]),
            Paragraph(report["title"], styles["Heading2"]),
            Paragraph(f"Generated {now().strftime('%d %b %Y %H:%M')}", styles["Normal"]),
            Spacer(1, 12),
        ]
        data = [report["headers"]]
        for row in report["rows"]:
            data.append([str(row.get(c) if row.get(c) is not None else "") for c in report["columns"]])
        if not report["rows"]:
            data.append(["No records"] + [""] * (len(report["headers"]) - 1))
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B2A4A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9A227")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F3E8")]),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 12))
        for k, v in report["summary"].items():
            story.append(Paragraph(f"<b>{k}</b>: {v}", styles["Normal"]))
        doc.build(story)
        self.audit.record(user, "reports.export", entity_type="report", new={"format": "pdf", "path": path.name})
        return path


def _slug(title: str) -> str:
    safe = "".join(ch if ch.isalnum() or ch in "-_ " else "" for ch in title)
    return safe.strip().replace(" ", "_")[:60] or "report"
